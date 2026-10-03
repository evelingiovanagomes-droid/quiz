"""Converte data/perguntas.xlsx (abas 'perguntas' e 'produtos') em CSV + JS."""
import csv, json, sys
from pathlib import Path
from openpyxl import load_workbook

base = Path(__file__).resolve().parent.parent / "data"
wb = load_workbook(base / "perguntas.xlsx", data_only=True)

def ler(nome):
    ws = wb[nome]
    ls = [[("" if c is None else str(c).strip()) for c in r] for r in ws.iter_rows(values_only=True)]
    return [l for l in ls if any(l)]

perguntas, produtos = ler("perguntas"), ler("produtos")
slugs = {l[0] for l in produtos[1:]}
erros = []
for n, l in enumerate(perguntas[1:], start=2):
    l += [""] * (10 - len(l))
    if not all(l[:9]): erros.append(f"Linha {n}: campo obrigatório vazio (id..explicacao_b, produtos)")
    elif l[5].upper() not in ("A", "B"): erros.append(f"Linha {n}: 'correta' deve ser A ou B")
    else:
        for s in l[8].split("|") + ([l[9]] if l[9] else []):
            if s not in slugs: erros.append(f"Linha {n}: produto '{s}' não existe na aba 'produtos'")
if erros:
    print("\n".join(erros)); sys.exit(1)

for nome, linhas, n in (("perguntas", perguntas, 10), ("produtos", produtos, 5)):
    with open(base / f"{nome}.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows([(l + [""] * n)[:n] for l in linhas])
print(f"OK: {len(perguntas)-1} perguntas e {len(produtos)-1} produtos exportados")
exec((Path(__file__).parent / "csv_para_js.py").read_text(encoding="utf-8"))
